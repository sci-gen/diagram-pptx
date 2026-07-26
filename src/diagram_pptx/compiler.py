"""High-level parsing, layout, style-resolution, and rendering compiler."""

from __future__ import annotations

from collections.abc import Mapping
from dataclasses import dataclass, field, replace
from enum import Enum
from typing import Any, Literal, TypedDict

from pptx.slide import Slide

from .diagnostics import (
    Diagnostic,
    MermaidRuntimeError,
    PartialModelMutationError,
)
from .layout.native import layout_native
from .mermaid import parse_mermaid, serialize_mermaid
from .model import MermaidDocument, MermaidSourceDiagram, SemanticDiagram
from .render.python_pptx import PythonPptxRenderer, RenderResult
from .scene import DrawingScene, SceneConnector, SceneContainer, SceneShape, SceneText
from .styles import (
    DEFAULT_THEME,
    ColorMapStyle,
    DiagramTheme,
    ElementStyle,
    SourceStylePolicy,
    StyleResolver,
    contrast_text_color,
    normalize_color,
    style_preset,
)


class CompileBackend(str, Enum):
    """Geometry backend used to create the renderer-neutral drawing scene."""

    AUTO = "auto"
    NATIVE = "native"
    OFFICIAL = "official"


BackendName = Literal["auto", "native", "official"]
StylePreset = Literal["native", "official"]
PlacementPreset = Literal["full", "left", "right", "top", "bottom"]
ColorMapName = Literal["jet", "viridis", "plasma", "magma"]
Bounds = tuple[float, float, float, float]
RelativeBounds = tuple[float, float, float, float]


class ColorMapOptions(TypedDict, total=False):
    """JSON-friendly continuous color-map configuration.

    ``name`` is required when this mapping is passed at runtime. Other values
    are sampling positions between ``0`` and ``1``.
    """

    name: ColorMapName
    primary: float
    secondary: float
    fill: float
    line: float
    text: float
    edge: float
    decision_fill: float
    group_line: float
    label_fill: float


_RELATIVE_PLACEMENTS: dict[str, tuple[float, float, float, float]] = {
    "full": (0.05, 0.05, 0.90, 0.90),
    "left": (0.05, 0.08, 0.42, 0.84),
    "right": (0.53, 0.08, 0.42, 0.84),
    "top": (0.05, 0.08, 0.90, 0.38),
    "bottom": (0.05, 0.54, 0.90, 0.38),
}


@dataclass(slots=True)
class CompileResult:
    """Result of compiling one diagram into a slide.

    Attributes:
        backend_used: Actual geometry backend, either ``"native"`` or
            ``"official"``.
        diagnostics: Non-fatal parse or runtime diagnostics.
        mermaid_version: Detected Mermaid CLI version for Official output.
        scene: Resolved, positioned drawing scene supplied to the renderer.
        render_result: Low-level collection of created ``python-pptx`` shapes.

    The convenience properties expose the most commonly needed native objects.
    ``group_shape`` is the outer editable group when ``group=True``.
    ``element_shapes`` maps semantic IDs to their primary child shapes.
    """

    backend_used: str
    diagnostics: list[Diagnostic] = field(default_factory=list)
    mermaid_version: str | None = None
    scene: DrawingScene | None = None
    render_result: RenderResult = field(default_factory=RenderResult)

    @property
    def group_shape(self) -> Any | None:
        return self.render_result.group_shape

    @property
    def element_shapes(self) -> dict[str, Any]:
        return self.render_result.element_shapes

    @property
    def element_parts(self) -> dict[str, list[Any]]:
        return self.render_result.element_parts

    @property
    def top_level_shapes(self) -> list[Any]:
        return self.render_result.top_level_shapes

    @property
    def node_shapes(self) -> dict[str, Any]:
        return self.render_result.node_shapes

    @property
    def connectors(self) -> list[Any]:
        return self.render_result.connectors

    @property
    def edge_label_shapes(self) -> list[Any]:
        return self.render_result.edge_label_shapes

    @property
    def group_shapes(self) -> list[Any]:
        return self.render_result.group_shapes

    @property
    def nested_group_shapes(self) -> list[Any]:
        return self.render_result.nested_group_shapes

    @property
    def shapes(self) -> list[Any]:
        return self.render_result.shapes

    def to_dict(self) -> dict[str, Any]:
        """Return a compact JSON-safe summary suitable for a tool response."""

        return {
            "backend_used": self.backend_used,
            "mermaid_version": self.mermaid_version,
            "diagram_kind": self.scene.kind if self.scene is not None else None,
            "diagnostics": [item.to_dict() for item in self.diagnostics],
            "grouped": self.group_shape is not None,
            "element_ids": list(self.element_shapes),
            "shape_count": len(self.shapes),
            "top_level_shape_count": len(self.top_level_shapes),
        }


@dataclass(slots=True)
class SceneResult:
    """Renderer-neutral result shared by PPTX and image exporters."""

    backend_used: str
    scene: DrawingScene
    diagnostics: list[Diagnostic] = field(default_factory=list)
    mermaid_version: str | None = None

    def to_dict(self) -> dict[str, Any]:
        """Return a compact JSON-safe summary suitable for a tool response."""

        return {
            "backend_used": self.backend_used,
            "mermaid_version": self.mermaid_version,
            "diagram_kind": self.scene.kind,
            "diagnostics": [item.to_dict() for item in self.diagnostics],
            "element_count": len(self.scene.elements),
            "width": self.scene.width,
            "height": self.scene.height,
        }


def build_scene(
    diagram: MermaidDocument | SemanticDiagram,
    *,
    backend: CompileBackend | BackendName = "native",
    style: StylePreset = "native",
    theme: DiagramTheme | Mapping[str, Any] | None = None,
    colors: ColorMapStyle | ColorMapOptions | ColorMapName | None = None,
    label_background: str | None = None,
    source_style: SourceStylePolicy = "merge",
    style_overrides: Mapping[str, ElementStyle | Mapping[str, Any]] | None = None,
    mmdc_path: str | None = None,
    strict: bool = False,
    timeout: float = 30.0,
    **legacy: Any,
) -> SceneResult:
    """Build a styled renderer-neutral scene without requiring a slide.

    This is the common compilation stage used by PowerPoint, SVG, PNG, and
    JPEG output. Most callers should use :func:`compile_diagram`,
    :func:`save_diagram`, or the convenience methods on
    :class:`MermaidDocument`.
    """

    colormap_style = _coerce_color_input(colors, legacy)
    if legacy:
        unknown = ", ".join(sorted(legacy))
        raise TypeError(f"Unexpected build_scene options: {unknown}")
    if timeout <= 0:
        raise ValueError("timeout must be greater than zero")
    document = (
        diagram
        if isinstance(diagram, MermaidDocument)
        else MermaidDocument(source=serialize_mermaid(diagram), model=diagram)
    )
    requested = CompileBackend(backend)
    if (
        not document.is_fully_modeled
        and document.model_changed
        and not isinstance(document.model, MermaidSourceDiagram)
    ):
        raise PartialModelMutationError(
            "This Mermaid document contains unsupported statements and its partial "
            "semantic model was changed. Compile the unchanged source with the official "
            "backend, or remove unsupported syntax before editing the model."
        )

    backend_used = requested
    mermaid_version: str | None = None
    diagnostics = list(document.diagnostics)

    if requested == CompileBackend.AUTO:
        from .official import find_mmdc

        executable = find_mmdc(mmdc_path)
        backend_used = CompileBackend.OFFICIAL if executable is not None else CompileBackend.NATIVE
        if executable is None:
            diagnostics.append(
                Diagnostic(
                    code="mmdc-not-found-native-fallback",
                    severity="info",
                    backend="native",
                    message=(
                        'mmdc was not found; backend="auto" selected the Native backend. '
                        "Install Node.js and @mermaid-js/mermaid-cli to enable Official "
                        "geometry."
                    ),
                )
            )

    if backend_used == CompileBackend.NATIVE:
        if isinstance(document.model, MermaidSourceDiagram):
            raise MermaidRuntimeError(
                f"Mermaid {document.model.kind!r} diagrams currently require the "
                "Official backend. Install Node.js, run `npm install -g "
                "@mermaid-js/mermaid-cli@11.16.0`, and use backend='official' "
                "or backend='auto'."
            )
        if not document.is_fully_modeled:
            raise MermaidRuntimeError(
                "The input contains Mermaid syntax that the native backend cannot model. "
                "Official geometry requires Node.js and @mermaid-js/mermaid-cli. Install "
                "Node.js, run `npm install -g "
                "@mermaid-js/mermaid-cli@11.16.0`, provide mmdc_path, or use strict "
                "parsing to locate unsupported statements."
            )
        scene = layout_native(document.model)
    else:
        from .official import render_official_scene

        source = serialize_mermaid(document.model) if document.model_changed else document.source
        official = render_official_scene(
            source,
            kind=document.model.kind,
            mmdc_path=mmdc_path,
            strict=strict,
            timeout=timeout,
        )
        scene = official.scene
        mermaid_version = official.version
        diagnostics.extend(official.diagnostics)

    resolved_theme = DEFAULT_THEME.merged(style_preset(style)).merged(theme)
    resolver = StyleResolver(resolved_theme, source_style=source_style)
    _resolve_scene_styles(
        scene,
        resolver,
        style_overrides or {},
        colormap_style,
        label_background=label_background,
    )
    return SceneResult(
        backend_used=backend_used.value,
        diagnostics=diagnostics,
        mermaid_version=mermaid_version,
        scene=scene,
    )


def compile_diagram(
    diagram: MermaidDocument | SemanticDiagram,
    *,
    slide: Slide,
    bounds: Bounds | None = None,
    position: PlacementPreset | None = None,
    relative_bounds: RelativeBounds | None = None,
    backend: CompileBackend | BackendName = "native",
    style: StylePreset = "native",
    theme: DiagramTheme | Mapping[str, Any] | None = None,
    colors: ColorMapStyle | ColorMapOptions | ColorMapName | None = None,
    label_background: str | None = None,
    source_style: SourceStylePolicy = "merge",
    style_overrides: Mapping[str, ElementStyle | Mapping[str, Any]] | None = None,
    group: bool = True,
    mmdc_path: str | None = None,
    strict: bool = False,
    timeout: float = 30.0,
    renderer: PythonPptxRenderer | None = None,
    **legacy: Any,
) -> CompileResult:
    """Compile a semantic diagram into an existing ``python-pptx`` slide.

    Args:
        diagram: Parsed :class:`MermaidDocument`, one of the five typed
            semantic models, or a lossless Official-only source model.
        slide: Existing ``python-pptx`` slide receiving the native shapes.
        bounds: Exact ``(left, top, width, height)`` in inches.
        position: Named slide-relative region. Supported values are ``full``,
            ``left``, ``right``, ``top``, and ``bottom``.
        relative_bounds: Normalized ``(x, y, width, height)`` using a top-left
            origin. Every value and both far edges must remain within ``0..1``.
        backend: Geometry source. ``native`` needs Python only; ``official``
            invokes Mermaid CLI; ``auto`` uses Official when available.
        style: Visual preset, independent of the geometry backend.
        theme: :class:`DiagramTheme` or an equivalent JSON-like mapping.
        colors: Continuous color-map settings. Pass a :class:`ColorMapStyle`,
            a map name, or a mapping such as
            ``{"name": "viridis", "primary": 0.8, "secondary": 0.2}``.
        label_background: Optional global connector/message-label background.
            Accepts RGB/RGBA, CSS colors, palette tokens, PowerPoint theme
            slots, or ``"transparent"``. Per-element overrides and a
            color-map ``label_fill`` channel take precedence.
        source_style: Precedence policy for Mermaid-authored styles.
        style_overrides: Final element styles keyed by semantic element ID.
        group: Put the complete diagram in one editable PowerPoint group.
        mmdc_path: Optional Mermaid CLI executable path.
        strict: Reject unsupported Mermaid syntax and untested Mermaid CLI
            versions instead of returning diagnostics.
        timeout: Maximum Mermaid CLI runtime in seconds.
        renderer: Advanced renderer injection point.

    Returns:
        A :class:`CompileResult` containing the scene, diagnostics, and native
        shape handles.

    Raises:
        ValueError: If placement, colors, backend, or styles are invalid.
        PartialModelMutationError: If an incomplete parsed model was mutated.
        MermaidRuntimeError: If the selected backend cannot render the input.

    Note:
        ``bounds``, ``position``, and ``relative_bounds`` are mutually
        exclusive. If all are omitted, ``position="full"`` is used.
    """

    scene_result = build_scene(
        diagram,
        backend=backend,
        style=style,
        theme=theme,
        colors=colors,
        label_background=label_background,
        source_style=source_style,
        style_overrides=style_overrides,
        mmdc_path=mmdc_path,
        strict=strict,
        timeout=timeout,
        **legacy,
    )
    resolved_bounds = resolve_diagram_bounds(
        slide,
        bounds=bounds,
        position=position,
        relative_bounds=relative_bounds,
    )
    native_result = (renderer or PythonPptxRenderer()).render(
        scene_result.scene,
        target=slide,
        bounds=resolved_bounds,
        group=group,
        group_name=f"diagram:{scene_result.scene.kind}",
    )
    return CompileResult(
        backend_used=scene_result.backend_used,
        diagnostics=scene_result.diagnostics,
        mermaid_version=scene_result.mermaid_version,
        scene=scene_result.scene,
        render_result=native_result,
    )


def render_mermaid(
    source: str,
    *,
    slide: Slide,
    bounds: Bounds | None = None,
    position: PlacementPreset | None = None,
    relative_bounds: RelativeBounds | None = None,
    backend: CompileBackend | BackendName = "native",
    style: StylePreset = "native",
    theme: DiagramTheme | Mapping[str, Any] | None = None,
    colors: ColorMapStyle | ColorMapOptions | ColorMapName | None = None,
    label_background: str | None = None,
    source_style: SourceStylePolicy = "merge",
    style_overrides: Mapping[str, ElementStyle | Mapping[str, Any]] | None = None,
    group: bool = True,
    mmdc_path: str | None = None,
    strict: bool = False,
    timeout: float = 30.0,
    renderer: PythonPptxRenderer | None = None,
    **legacy: Any,
) -> CompileResult:
    """Parse Mermaid and render editable shapes into an existing slide.

    This is the recommended one-call API. Its options are identical to
    :func:`compile_diagram`, except that ``source`` is parsed first.

    Args:
        source: Complete Mermaid source for one supported diagram.
        slide: Existing ``python-pptx`` slide receiving the diagram.
        bounds: Exact ``(left, top, width, height)`` in inches.
        position: ``full``, ``left``, ``right``, ``top``, or ``bottom``.
        relative_bounds: Normalized top-left-origin
            ``(x, y, width, height)`` inside ``0..1``.
        backend: ``native`` (default), ``official``, or ``auto``.
        style: ``native`` or Mermaid-like ``official`` visual preset.
        theme: Theme object or JSON-like theme mapping.
        colors: Color-map name, :class:`ColorMapStyle`, or JSON-like mapping.
        label_background: Optional background behind connector and message
            labels. Use an explicit slide/canvas color or ``"transparent"``.
        source_style: Mermaid style precedence policy.
        style_overrides: Element styles keyed by semantic ID.
        group: Group all generated shapes as one editable object.
        mmdc_path: Optional Mermaid CLI executable path.
        strict: Reject unsupported syntax and untested Mermaid CLI versions.
        timeout: Mermaid CLI timeout in seconds.
        renderer: Advanced renderer injection point.

    Returns:
        A :class:`CompileResult` with native shape handles and diagnostics.

    Examples:
        Place two independent diagrams on one slide::

            render_mermaid(left_source, slide=slide, position="left")
            render_mermaid(right_source, slide=slide, position="right")

        Use JSON-friendly styling::

            render_mermaid(
                source,
                slide=slide,
                relative_bounds=(0.05, 0.1, 0.4, 0.8),
                style="official",
                colors={
                    "name": "viridis",
                    "primary": 0.82,
                    "secondary": 0.18,
                },
            )
    """

    # Keep the old injectable flow components usable for one alpha transition.
    importer = legacy.pop("importer", None)
    layout_engine = legacy.pop("layout_engine", None)
    colormap_style = _coerce_color_input(colors, legacy)
    if legacy:
        unknown = ", ".join(sorted(legacy))
        raise TypeError(f"Unexpected render_mermaid options: {unknown}")
    if importer is not None or layout_engine is not None:
        if backend != CompileBackend.NATIVE:
            raise TypeError("Legacy importer/layout_engine cannot be combined with backend")
        model = importer.parse(source) if importer else parse_mermaid(source, strict=strict).model
        if layout_engine is not None:
            # Custom legacy layouts are adapted through the flow scene conversion path.
            from .layout.native import _layout_flow

            if layout_engine.__class__.__name__ != "LayeredLayout":
                raise TypeError(
                    "Custom legacy layout engines must migrate to a DrawingScene backend"
                )
            scene = _layout_flow(model)
            resolved_theme = DEFAULT_THEME.merged(style_preset(style)).merged(theme)
            resolver = StyleResolver(resolved_theme, source_style=source_style)
            _resolve_scene_styles(
                scene,
                resolver,
                style_overrides or {},
                colormap_style,
                label_background=label_background,
            )
            native_result = (renderer or PythonPptxRenderer()).render(
                scene,
                target=slide,
                bounds=resolve_diagram_bounds(
                    slide,
                    bounds=bounds,
                    position=position,
                    relative_bounds=relative_bounds,
                ),
                group=group,
            )
            return CompileResult(
                backend_used="native",
                scene=scene,
                render_result=native_result,
            )
        return compile_diagram(
            model,
            slide=slide,
            bounds=bounds,
            position=position,
            relative_bounds=relative_bounds,
            backend="native",
            style=style,
            theme=theme,
            colors=colormap_style,
            label_background=label_background,
            source_style=source_style,
            style_overrides=style_overrides,
            group=group,
            strict=strict,
            renderer=renderer,
        )

    return compile_diagram(
        parse_mermaid(source, strict=strict),
        slide=slide,
        bounds=bounds,
        position=position,
        relative_bounds=relative_bounds,
        backend=backend,
        style=style,
        theme=theme,
        colors=colormap_style,
        label_background=label_background,
        source_style=source_style,
        style_overrides=style_overrides,
        group=group,
        mmdc_path=mmdc_path,
        strict=strict,
        timeout=timeout,
        renderer=renderer,
    )


def _resolve_scene_styles(
    scene: DrawingScene,
    resolver: StyleResolver,
    overrides: Mapping[str, ElementStyle | Mapping[str, Any]],
    colormap: ColorMapStyle | None = None,
    *,
    label_background: str | None = None,
) -> None:
    for element in scene.elements:
        if isinstance(element, SceneConnector):
            base_role = "edge.default"
        elif isinstance(element, SceneContainer):
            base_role = "group.default"
        elif isinstance(element, SceneText):
            base_role = "text.default"
        elif isinstance(element, SceneShape):
            base_role = "node.default"
        else:
            continue
        base = resolver.theme.roles.get(base_role, ElementStyle())
        exact = resolver.theme.roles.get(element.role)
        global_override = (
            ElementStyle(label_fill=label_background)
            if label_background is not None
            and (
                isinstance(element, SceneConnector)
                and bool(element.label)
                or isinstance(element, SceneText)
                and element.role == "edge.label"
            )
            else ElementStyle()
        )
        explicit_override = ElementStyle.from_dict(overrides.get(element.semantic_id))
        element_override = global_override.merged(explicit_override)
        element.style = resolver.resolve(
            element_id=element.semantic_id,
            role=element.role,
            classes=element.classes,
            source=element.style,
            default=base.merged(exact),
            override=element_override,
        )
        if (
            label_background is not None
            and explicit_override.text is None
            and element.style.label_fill is not None
        ):
            normalized_label_fill = normalize_color(element.style.label_fill)
            if normalized_label_fill.startswith("#") and (
                len(normalized_label_fill) == 7 or normalized_label_fill[7:9] != "00"
            ):
                element.style.text = contrast_text_color(normalized_label_fill)
        if colormap is not None:
            _apply_colormap_style(element, colormap)


def _coerce_colormap_style(
    colormap: str | ColorMapStyle | None,
    **positions: float | None,
) -> ColorMapStyle | None:
    configured = {name: value for name, value in positions.items() if value is not None}
    if colormap is None:
        if configured:
            raise ValueError("Color positions require a colormap name")
        return None
    if isinstance(colormap, ColorMapStyle):
        return replace(colormap, **configured) if configured else colormap
    return ColorMapStyle(name=colormap, **configured)


def _coerce_color_input(
    colors: ColorMapStyle | ColorMapOptions | ColorMapName | None,
    legacy: dict[str, Any],
) -> ColorMapStyle | None:
    """Normalize the canonical color object and the alpha compatibility form."""

    legacy_colormap = legacy.pop("colormap", None)
    legacy_positions = {
        name: legacy.pop(name, None)
        for name in (
            "primary",
            "secondary",
            "fill",
            "line",
            "text",
            "edge",
            "decision_fill",
            "group_line",
            "label_fill",
        )
    }
    legacy_is_configured = legacy_colormap is not None or any(
        value is not None for value in legacy_positions.values()
    )
    if colors is not None and legacy_is_configured:
        raise TypeError("Use colors={...}; do not combine it with legacy colormap/channel options")
    if colors is not None:
        if isinstance(colors, ColorMapStyle):
            return colors
        if isinstance(colors, Mapping):
            return ColorMapStyle.from_dict(colors)
        return ColorMapStyle(name=colors)
    return _coerce_colormap_style(legacy_colormap, **legacy_positions)


def _apply_colormap_style(element: Any, colormap: ColorMapStyle) -> None:
    explicit_text = colormap.color("text")
    if isinstance(element, SceneConnector):
        element.style.line = colormap.color("edge")
        if explicit_text is not None:
            element.style.text = explicit_text
        label_fill = colormap.color("label_fill")
        if label_fill is not None:
            element.style.label_fill = label_fill
        return
    if isinstance(element, SceneContainer):
        element.style.line = colormap.color("group_line") or colormap.color("line")
        if explicit_text is not None:
            element.style.text = explicit_text
        return
    if isinstance(element, SceneText):
        if explicit_text is not None:
            element.style.text = explicit_text
        return
    if isinstance(element, SceneShape):
        fill_channel = (
            "decision_fill" if element.shape == "diamond" or "decision" in element.role else "fill"
        )
        fill_color = colormap.color(fill_channel)
        element.style.fill = fill_color
        element.style.line = colormap.color("line")
        element.style.text = (
            explicit_text
            if explicit_text is not None
            else contrast_text_color(fill_color or "#FFFFFF")
        )


def resolve_diagram_bounds(
    slide: Slide,
    *,
    bounds: Bounds | None = None,
    position: PlacementPreset | None = None,
    relative_bounds: RelativeBounds | None = None,
) -> Bounds:
    """Resolve one placement option to absolute slide coordinates in inches.

    Args:
        slide: Existing ``python-pptx`` slide whose presentation dimensions
            define normalized coordinates.
        bounds: Exact ``(left, top, width, height)`` in inches.
        position: Named preset: ``full``, ``left``, ``right``, ``top``, or
            ``bottom``.
        relative_bounds: Normalized top-left-origin
            ``(x, y, width, height)`` inside ``0..1``.

    Returns:
        ``(left, top, width, height)`` in inches.

    Raises:
        ValueError: If multiple placement options are supplied, dimensions are
            non-positive, or a normalized rectangle exceeds the slide.

    If no placement option is supplied, the ``full`` preset is used.
    """

    configured = sum(value is not None for value in (bounds, position, relative_bounds))
    if configured > 1:
        raise ValueError("Use only one of bounds, position, or relative_bounds")
    if bounds is not None:
        left, top, width, height = (float(value) for value in bounds)
        if left < 0 or top < 0 or width <= 0 or height <= 0:
            raise ValueError("bounds must be left,top,positive-width,positive-height")
        return left, top, width, height

    if position is not None:
        normalized_position = str(position).strip().lower()
        if normalized_position not in _RELATIVE_PLACEMENTS:
            supported = ", ".join(_RELATIVE_PLACEMENTS)
            raise ValueError(f"Unknown position {position!r}; choose one of: {supported}")
        relative = _RELATIVE_PLACEMENTS[normalized_position]
    else:
        relative = relative_bounds or _RELATIVE_PLACEMENTS["full"]
    x, y, width, height = (float(value) for value in relative)
    if (
        not 0 <= x <= 1
        or not 0 <= y <= 1
        or width <= 0
        or height <= 0
        or x + width > 1
        or y + height > 1
    ):
        raise ValueError("relative_bounds must be x,y,positive-width,positive-height inside 0..1")
    presentation = slide.part.package.presentation_part.presentation
    slide_width = presentation.slide_width / 914400
    slide_height = presentation.slide_height / 914400
    return (
        x * slide_width,
        y * slide_height,
        width * slide_width,
        height * slide_height,
    )
